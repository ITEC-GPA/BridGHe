"""Create the Italian BridGHe operator guide using python-docx and Pillow.

Run with the Codex bundled Python runtime (python-docx and Pillow required).
The DOCX is the deliverable; intermediate figures remain under artifacts/.
After changing this script, render the DOCX and inspect every page before release.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "BridGHe_Guida_operativa.docx"
FIGURES = ROOT / "artifacts" / "document-guide" / "figures"
WIDTH = 6.94
BLUE = "24465F"


def font(size, bold=False):
    return ImageFont.truetype(
        "C:/Windows/Fonts/" + ("arialbd.ttf" if bold else "arial.ttf"), size
    )


def figures():
    FIGURES.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (1800, 640), "white")
    d = ImageDraw.Draw(im)
    rows = [
        ("1  Modello acciaio", "2  Proprietà per fase", "3  Analisi Straus"),
        ("4  Azioni e ritiro", "5  Fasi e combinazioni", "6  Verifiche e grafici"),
    ]
    subtitles = [
        ("Geometria e ID coerenti", "Override dei longitudinali", "Incrementi separati"),
        ("Sorgente e strain associati", "Numeri e Action ID espliciti", "Risultati tracciabili"),
    ]
    for row in range(2):
        y = 110 + row * 290
        for col in range(3):
            x = 25 + col * 615
            d.text((x, y), rows[row][col], font=font(38, True), fill="black")
            d.text((x, y + 65), subtitles[row][col], font=font(30), fill="#454545")
            if col < 2:
                ax = x + 530
                d.line((ax, y + 25, ax + 55, y + 25), fill=f"#{BLUE}", width=4)
                d.polygon([(ax + 55, y + 25), (ax + 42, y + 15), (ax + 42, y + 35)], fill=f"#{BLUE}")
    d.line((1560, 230, 1560, 305, 230, 305, 230, 365), fill=f"#{BLUE}", width=4)
    d.polygon([(230, 377), (219, 360), (241, 360)], fill=f"#{BLUE}")
    im.save(FIGURES / "workflow.png")

    im = Image.new("RGB", (1800, 730), "white")
    d = ImageDraw.Draw(im)
    x0, x1 = 150, 1700
    ys = [170, 330, 490]
    for k in range(13):
        x = x0 + (x1 - x0) * k / 12
        d.line((x, ys[0], x, ys[2]), fill="#A6ACB1", width=4)
    for j, y in enumerate(ys):
        d.line((x0, y, x1, y), fill=f"#{BLUE}", width=8 if j == 1 else 5)
        d.text((5, y - 20), ["G3", "G2", "G1"][j], font=font(34, True), fill="black")
        for k in range(25):
            x = x0 + (x1 - x0) * k / 24
            d.ellipse((x - 6, y - 6, x + 6, y + 6), fill="black")
        for k in [0, 12, 24]:
            x = x0 + (x1 - x0) * k / 24
            d.polygon([(x, y + 8), (x - 17, y + 34), (x + 17, y + 34)], outline="black", width=3)
    d.text((150, 50), "Schema in pianta  |  esempio 19 REV003", font=font(36, True), fill="black")
    for x, label in [(x0, "x = 0 m"), ((x0 + x1) / 2, "x = 30 m"), (x1 - 150, "x = 60 m")]:
        d.text((x, 565), label, font=font(30), fill="black")
    d.text((150, 635), "Longitudinali ogni 4 m  •  traversi ogni 5 m  •  triangoli = appoggi verticali", font=font(30), fill="black")
    im.save(FIGURES / "bridge-plan.png")


def make_document():
    figures()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.73), Inches(.69)
    sec.left_margin = sec.right_margin = Inches(.78)
    sec.header_distance = sec.footer_distance = Inches(.30)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.06
    for name, size in [("Title", 27), ("Subtitle", 15), ("Heading 1", 19), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name.startswith("Heading")
        style.paragraph_format.space_before = Pt(12 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
        for node in list(style.element.xpath("./w:pPr/w:pBdr")):
            node.getparent().remove(node)
    doc.styles["Caption"].font.name = "Calibri"
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("444444")
    doc.styles["Caption"].font.italic = False
    sec.header.paragraphs[0].text = "BridGHe   |   Guida operativa"
    sec.header.paragraphs[0].runs[0].font.size = Pt(9)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run("20 ottobre 2026   •   ").font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    foot._p.append(fld)
    props = doc.core_properties
    props.title = "Guida operativa BridGHe"
    props.subject = "Workflow acciaio misto con fasi numerate e azioni incrementali"
    props.author = "BridGHe"
    props.keywords = "BridGHe, Grasshopper, Rhino 8, Straus7 R3, fasi, ritiro"
    props.comments = "Workflow REV003, esempi numerati e strumenti operativi 23–28."

    def p(text, style=None):
        par = doc.add_paragraph(style=style)
        # Lightweight emphasis, avoiding code fonts in ordinary component names.
        for i, part in enumerate(re.split(r"\*\*(.*?)\*\*", text)):
            r = par.add_run(part)
            r.bold = bool(i % 2)
        return par

    def h(text):
        doc.add_heading(text, 2)

    def page(number, title):
        doc.add_page_break()
        doc.add_heading(f"{number}  {title}", 1)

    def steps(items):
        for i, item in enumerate(items, 1):
            par = p(f"**{i}.** {item}")
            par.paragraph_format.left_indent = Inches(.18)
            par.paragraph_format.first_line_indent = Inches(-.18)

    def table(headers, rows, widths, centers=()):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        pr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement("w:" + side)
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                border.set(qn("w:" + key), value)
            borders.append(border)
        pr.append(borders)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row_idx, values in enumerate([headers] + list(rows)):
            row = t.rows[0] if row_idx == 0 else t.add_row()
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for side, val in [("top", "75"), ("bottom", "75"), ("left", "100"), ("right", "100")]:
                    el = OxmlElement("w:" + side)
                    el.set(qn("w:w"), val)
                    el.set(qn("w:type"), "dxa")
                    mar.append(el)
                cp.append(mar)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), BLUE if row_idx == 0 else ("F2F5F7" if row_idx % 2 == 0 else "FFFFFF"))
                cp.append(shade)
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.03
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in centers else WD_ALIGN_PARAGRAPH.LEFT
                run = par.add_run(str(value))
                run.font.size = Pt(10)
                run.bold = row_idx == 0
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_idx == 0 else "000000")
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(1)
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.line_spacing = 1
        after.add_run().font.size = Pt(2)
        return t

    def code(lines):
        for line in lines.splitlines():
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)

    # Page 1
    doc.add_paragraph("Guida operativa BridGHe", "Title")
    doc.add_paragraph("Dal modello parametrico alla verifica per fasi", "Subtitle")
    p("Rhino 8 • Grasshopper • Straus7 R3 • librerie gPC")
    p("Questa guida spiega come costruire un modello in acciaio, analizzare gli incrementi di carico con le proprietà delle diverse fasi del misto, associare ritiro e risultati FEM e assemblare il verificatore. Il riferimento operativo è il ponte a due campate e tre travi dell’**esempio 19 REV003**.")
    p("Le fasi sono ordinate tramite numeri espliciti; ogni azione conserva il proprio modello analizzato e il caso sorgente. Le combinazioni utilizzano gli **Action ID**. Questo collegamento rende identificabile l’origine di ogni contributo senza dipendere dall’ordine dei modelli.")
    doc.add_picture(str(FIGURES / "workflow.png"), width=Inches(WIDTH))
    h("Percorso di lettura")
    table(["Capitoli", "Contenuto", "Pagine"], [
        ("1–4", "Ambiente, sezioni, modello acciaio e override", "2–5"),
        ("5–9", "Ritiro, analisi, azioni, fasi e combinazioni", "6–10"),
        ("10–12", "Verificatore, dettagli e visualizzazione", "11–13"),
        ("13–14", "Carichi mobili ed esecuzione dell’esempio 19", "14–15"),
        ("15–16", "Diagnostica, prove e riferimenti riproducibili", "16–17"),
        ("17–21", "Nuovi strumenti operativi ed esempi 23–28", "18–22"),
        ("22–27", "Dal tracciato al modello ed esempi 29–45", "23–28"),
    ], [1.0, 5.15, .79], centers=(0, 2))
    p("**Edizione aggiornata:** 20 ottobre 2026, workflow REV003, Simple Bridge e strumenti di modellazione 33–45. Le interfacce precedenti di Build BridGHe Model non sono compatibili con il workflow corrente. La guida distingue controlli offline e analisi Straus.")
    p("Il percorso produce controlli di sezione nel campo implementato. L’esempio usa azioni e coefficienti didattici; le verifiche di ponte ancora necessarie sono indicate nei capitoli 10, 11 e 14.")

    # Page 2
    page(1, "Preparare ambiente e unità")
    steps([
        "Utilizzare **Rhino 8 su Windows x64 con .NET 8**. Per l’analisi deve essere disponibile l’installazione Straus7 R3 e la relativa licenza API.",
        "Installare **BridGHe.Grasshopper.gha e tutte le DLL associate** dalla cartella bin/Release, mantenendole insieme nella cartella Libraries di Grasshopper o nel percorso di caricamento configurato. Rimuovere dal caricamento le copie obsolete del plugin.",
        "Prima di sostituire DLL già caricate, salvare e chiudere Rhino. Riavviare Rhino e Grasshopper e verificare la presenza dell’unica tab **BridGHe**.",
        "Verificare il percorso di Straus. Il valore predefinito è C:\\Program Files\\Straus7 R31; per un’altra installazione impostare BRIDGHE_STRAUS_HOME prima di avviare Rhino.",
        "Aprire una copia di lavoro di 19-flusso-acciaio-misto_REV003.ghx. Lasciare inizialmente disattivato CALCOLA 4 MODELLI STRAUS e controllare i percorsi di output .st7.",
    ])
    h("Sistema di unità del workflow")
    table(["Grandezza", "Unità e convenzione"], [
        ("Coordinate e dimensioni", "mm; una campata da 30 m si inserisce come 30000"),
        ("Forze e carichi lineari", "N e N/mm; 1 N/mm = 1 kN/m"),
        ("Momenti e proprietà", "N·mm, area mm², inerzie e J in mm⁴"),
        ("Materiali e tensioni", "MPa = N/mm²; rapporti n adimensionali"),
        ("Strain imposto", "Adimensionale; −200 µε = −0.0002"),
        ("Grafici delle azioni", "kN e kNm; progressive in m"),
    ], [2.0, 4.94])
    p("Impostare **N–mm** nel modello. Non è prevista una conversione implicita delle coordinate da metri a millimetri. Se la geometria proviene da un documento Rhino in metri, convertirla prima di creare gli elementi.")
    h("File e organizzazione")
    p("Ogni esempio è in example, in una cartella numerata con README e, quando prevista, definizione GHX. La sottocartella calcoli contiene modelli .st7 e risultati; archivio conserva le revisioni precedenti. Usare file diversi per ogni fase. L’indice completo è example/README.md; bin/Release contiene il plugin installabile.")
    p("Gli script di generazione possono sovrascrivere gli esempi generati: eseguirli dopo aver salvato le proprie modifiche con un nome distinto. REV001 e REV002 conservati dall’utente sono riferimenti storici, non la definizione corrente.")

    # Page 3
    page(2, "Definire la sezione mista")
    steps([
        "Creare la sezione di acciaio con **I Section**. Inserire larghezza della flangia inferiore, larghezza della superiore, altezza totale, spessore inferiore, spessore superiore e spessore dell’anima nell’ordine dei rispettivi ingressi. Non scambiare le due flange.",
        "Creare **Concrete Slab Section** con Height e Width in mm. Width è la larghezza efficace assegnata al modello: il componente non la determina automaticamente da una norma.",
        "Impostare Rebar area e Bottom rebar area come aree totali in mm² sull’intera larghezza assegnata. Top concrete cover e Bottom concrete cover sono distanze dalle rispettive facce all’asse delle barre. Area zero disattiva la fila.",
        "Collegare geometrie e materiali a **Composite Beam Property**. Reinforcement Material è necessario se sono presenti armature attive. Impostare Consider Concrete e Consider Rebars secondo la sezione fisica da considerare.",
        "Assegnare n0, n ∞ e n ritiro. Se n ritiro è omesso, viene usato n ∞. Nell’esempio i valori sono rispettivamente 7, 18 e 12; non costituiscono una scelta automatica di progetto.",
    ])
    table(["Uscita di Composite Beam Property", "Destinazione e significato"], [
        ("Composite Section", "Composite beam → Section Start; descrizione fisica per la verifica"),
        ("Only Steel Property", "Beam del modello acciaio iniziale"),
        ("n0 Beam Section", "Beam Override → Property per azioni di breve durata"),
        ("n ∞ Beam Section", "Beam Override → Property per azioni di lunga durata"),
        ("n ritiro Beam Section", "Beam Override → Property per il modello degli effetti secondari di ritiro"),
    ], [2.45, 4.49])
    p("Le proprietà equivalenti esportate in Straus sono sezioni generiche con area e inerzie omogeneizzate. Il rapporto n = Eₛ/Ec,eff è adimensionale. La rigidezza torsionale usa **J dell’acciaio**: non è una valutazione automatica della torsione della sezione mista.")
    h("Controllo prima di proseguire")
    p("Visualizzare la sezione e verificare altezza totale, posizione della soletta e delle due file di armatura. Accertare che area e inerzie siano positive e plausibili. La preview delle armature usa tre linee schematiche per fila attiva e non rappresenta una distinta ferri.")
    p("La sezione fisica rimane associata alle proprietà generiche: dopo Beam Override, la preview può ancora mostrare acciaio e soletta. Le proprietà native Straus possono richiedere la licenza API anche per il solo calcolo geometrico.")

    # Page 4
    page(3, "Assemblare il modello in acciaio")
    steps([
        "Definire gli assi di longitudinali e traversi con nodi comuni alle intersezioni. Suddividere le aste in corrispondenza di appoggi, variazioni geometriche e punti in cui occorre leggere i risultati.",
        "Creare i Beam con Only Steel Property. Per ogni famiglia assegnare materiale, orientamento e attributi. Controllare il verso degli assi locali: influisce sui segni di carichi, tagli e momenti.",
        "Assegnare ID univoci e stabili con **Element Id** e una famiglia con **Super Element**, per esempio G1, G2, G3 e CROSS. Conservare gli stessi ID in tutti i modelli di fase.",
        "Creare i nodi con Node e i vincoli con **Restraints**. Bloccare i moti rigidi lasciando liberi i gradi di libertà previsti dallo schema statico; non correggere una labilità con vincoli aggiunti senza giustificazione fisica.",
        "Raccogliere nodi, beam e attributi in **Build Model**, selezionando N–mm. Appiattire la lista degli elementi solo dove serve raccogliere le famiglie in un unico modello.",
    ])
    doc.add_picture(str(FIGURES / "bridge-plan.png"), width=Inches(WIDTH))
    p("Schema dell’esempio 19 REV003. Tre assi longitudinali a y = −4, 0, +4 m; due campate da 30 m. La larghezza nominale dell’impalcato è 12 m.", "Caption")
    h("Vincoli e discretizzazione di REV003")
    p("Tz è bloccata a x = 0, 30 e 60 m su ciascuna delle tre travi. Tx è bloccata nel primo appoggio esterno; Ty nel primo e nell’ultimo appoggio della stessa trave. Le rotazioni sono libere. La continuità delle aste sulla pila è mantenuta.")
    p("Ogni longitudinale ha 24 elementi da 2500 mm. Tredici allineamenti di traversi collegano le tre travi: 26 traversi, 72 elementi longitudinali, **98 beam e 75 nodi** complessivi.")
    p("**Controllo:** verificare connessioni, duplicati, appoggi, ID e assi nella preview. Build Model assembla il FEM; Build BridGHe Model, usato più avanti, assembla invece il modello del verificatore.")

    # Page 5
    page(4, "Creare le copie di fase tramite override")
    steps([
        "Partire dal modello acciaio e usare Decompose Model per recuperarne nodi ed elementi. Selezionare separatamente i longitudinali centrali, quelli laterali e i traversi.",
        "Applicare **Beam Override** ai soli longitudinali. Collegare a Property l’uscita omogeneizzata corrispondente alla fase e alla geometria di ciascuna famiglia.",
        "Ricostruire ogni copia con gli stessi nodi, supporti, traversi, ID, assi e offset. I traversi di REV003 mantengono la sezione di acciaio in tutte le fasi.",
        "Creare un **Load Case** per ogni incremento e collegarlo a **Beam Distributed Load** insieme ai beam. In REV003: Type = 0, Coordinate System locale, Direction = 1 (asse locale 2), PA = −33 oppure −10 N/mm e a = b = 0. Con gli assi dell’esempio questo carico è verticale verso il basso.",
        "Collegare a Build Model l’uscita dei beam caricati, insieme ai nodi e ai traversi. Non aggiungere anche le copie originali degli stessi beam. Non includere carichi già conteggiati nelle altre fasi.",
    ])
    table(["Fase", "Azione", "Proprietà e caso Straus", "Carico REV003"], [
        ("10", "G1", "Only Steel Property\nG1_GETTO", "33 N/mm verso il basso"),
        ("20", "G2", "n ∞ Beam Section\nG2_PERMANENTI", "10 N/mm verso il basso"),
        ("30", "Q", "n0 Beam Section\nQ_TRAFFICO", "10 N/mm verso il basso"),
        ("40", "RITIRO", "n ritiro Beam Section\nRITIRO_FEM", "Prestrain e gradiente equivalenti"),
    ], [.5, .75, 3.10, 2.59], centers=(0, 1))
    p("I carichi lineari sono applicati a ciascun longitudinale. In questo esempio non è attivo il peso proprio automatico. La corrispondenza fra direzione di carico e verticale deve essere verificata sul modello e nei risultati Straus.")
    h("Perché gli incrementi devono restare separati")
    p("G1 produce tensioni nell’acciaio prima dell’attivazione della soletta. G2 agisce sulla sezione mista con n ∞ e Q sulla sezione mista con n0. Sommare in seguito i contributi mantiene questa distinzione. Se il risultato G2 includesse già G1, la combinazione G1 + G2 conterebbe due volte G1.")
    h("Geometria comune e risultati confrontabili")
    p("Gli stessi Beam ID devono identificare gli stessi elementi fisici in ogni copia: estremi, orientamento, offset e unità devono essere coerenti. Anche le stazioni di risultato devono essere compatibili. Il workflow non interpola implicitamente attraverso campioni mancanti o salti di sollecitazione.")
    p("**Controllo:** confrontare conteggi e ID delle copie. La modifica intenzionale deve riguardare le proprietà e le azioni previste; modifiche topologiche richiedono un workflow specifico e non sono simulate dal solo numero di fase.")

    # Page 6
    page(5, "Modellare il ritiro senza doppio conteggio")
    p("Il ritiro della soletta produce tensioni primarie nella sezione composta e può generare azioni secondarie per effetto dei vincoli globali. **Bridge Shrinkage Action** associa i due contributi alla stessa azione e allo stesso coefficiente di combinazione.")
    steps([
        "Inserire lo strain del calcestruzzo, negativo per contrazione. Per −200 µε usare **−0.0002**. Collegarlo a Concrete strain di Composite Shrinkage Actions e a Strain increment di Bridge Shrinkage Action.",
        "Collegare Composite Section a **Composite Shrinkage Actions → Composite property**. Per ciascuna famiglia usare la propria sezione e lo stesso n ritiro impiegato nel modello di fase.",
        "Collegare Straus pre-strain a **Beam PreLoad → Value**, con Type = 1 (Strain). Passare il beam risultante a **Beam Gradient Load** e collegare il gradiente a Axis 2. Entrambi usano il caso RITIRO_FEM; l’uscita dell’ultimo componente entra in Build Model.",
        "Analizzare il modello con n ritiro Beam Section. Collegare il modello analizzato e RITIRO_FEM a Bridge Shrinkage Action. Lasciare Primary Only = false e Response = 0 per il ritiro ordinario.",
    ])
    table(["Uscita del generatore", "Significato e uso"], [
        ("Straus pre-strain", "Convenzione Straus: positivo per la contrazione equivalente. Collegamento diretto al preload."),
        ("Axis 2 thermal gradient", "Gradiente equivalente per unità di lunghezza; richiede il coefficiente di dilatazione dell’acciaio."),
        ("Free centroid strain", "Deformazione libera del baricentro omogeneizzato; controllo, non sostituisce εcs."),
        ("Free curvature", "Curvatura libera in 1/mm; controllo di ordine di grandezza e segno."),
    ], [2.12, 4.82])
    h("Significato del calcolo equivalente")
    p("Con Ec,eff = Eₛ/nritiro, il generatore usa N* = Ec,eff Ac εcs e M* = Ec,eff (Sc − Ac yG) εcs. Ac è l’area netta di calcestruzzo, Sc il suo momento statico e yG il baricentro omogeneizzato. Ne derivano ε0 = N*/(Eₛ Aeq) e κ = M*/(Eₛ Ieq), con Aeq e Ieq omogeneizzate all’acciaio. Infine prestrain = −ε0 e gradiente = κ/αₛ. N* è in N e M* in N·mm.")
    p("Non invertire manualmente i segni generati. La direzione di curvatura deve essere compatibile con gli assi locali; configurazioni inclinate non supportate sono rifiutate. Gli incrementi −0.00012 e −0.00008 producono complessivamente −0.00020.")
    p("**Primary Only = true** omette il secondario FEM: usarlo solo se questa è l’ipotesi voluta e motivata. Non aggiungere un ulteriore Composite Beam Strain alla stessa azione, perché duplicherebbe il contributo primario.")

    # Page 7
    page(6, "Eseguire le analisi Straus")
    steps([
        "Per ogni copia inserire **Straus R3 Analysis**. Collegare il modello FEM a Model e impostare Solver = 0 per l’analisi statica lineare adottata nell’esempio.",
        "Assegnare Output Path a un file .st7 distinto, in una cartella scrivibile. I quattro casi non devono sovrascriversi a vicenda.",
        "In Analysis Results selezionare il caso o i casi da leggere. Una Value List collegata può mostrare i nomi disponibili; verificare sempre che il caso necessario all’azione sia presente nei risultati.",
        "Avviare con il pulsante Run. Keep Running consente l’esecuzione a ogni soluzione, ma conviene attivarlo solo durante il calcolo voluto e poi riportarlo a false.",
        "Attendere la conclusione e controllare i messaggi. L’uscita Model è il modello con le sollecitazioni memorizzate negli elementi; è questa l’uscita da collegare alle azioni, non il modello FEM privo di risultati.",
    ])
    h("Controlli prima del verificatore")
    table(["Controllo", "Esito richiesto"], [
        ("Convergenza e schema", "Nessuna labilità inattesa; spostamenti e rotazioni plausibili"),
        ("Equilibrio", "Somma delle reazioni coerente con le azioni applicate"),
        ("Caso sorgente", "Nome esatto presente nel modello analizzato"),
        ("Campionamento", "Risultati finiti per ogni beam e stazione richiesti"),
        ("Segni e unità", "N, mm, MPa e assi locali coerenti fra le quattro copie"),
    ], [1.62, 5.32])
    p("Per REV003 i carichi distribuiti complessivi sono 5940 kN per G1 e 1800 kN per G2 e Q, separatamente. Sono controlli di equilibrio globale, non valori delle singole reazioni. Il ritiro è una deformazione imposta: le reazioni possono essere non nulle ai singoli appoggi pur essendo globalmente autoequilibrate.")
    h("Aggiornamento e cache")
    p("Il componente conserva il modello analizzato fra le soluzioni. Dopo una variazione di carichi, sezioni, vincoli o geometria **rieseguire tutte le analisi interessate** prima di aggiornare il verificatore. Una preview aggiornata non dimostra che anche i risultati FEM siano stati ricalcolati.")
    p("Il comando contestuale Reset cancella la cache. Se un’analisi fallisce, controllare l’errore e ottenere nuovamente un modello valido prima di proseguire; non utilizzare file storici della vecchia trave come risultati del nuovo graticcio.")
    p("In REV003 il controllo CALCOLA 4 MODELLI STRAUS comanda le quattro analisi. Dopo la conclusione riportarlo a false, quindi eseguire Build BridGHe Model e Super Element Check nell’ordine indicato al capitolo 14.")

    # Page 8
    page(7, "Creare le azioni del verificatore")
    p("Un’azione descrive un incremento e ne conserva la provenienza. Il suo **Action ID** è la chiave da riportare nella tabella delle combinazioni; il nome Case identifica invece il caso realmente calcolato in Straus. I due nomi possono essere diversi.")
    h("Bridge Action per carichi meccanici")
    table(["Ingresso", "Impostazione"], [
        ("Action ID", "Chiave stabile e univoca, per esempio G1, G2 o Q"),
        ("Model", "Uscita del corrispondente Straus R3 Analysis"),
        ("Case", "Oggetto caso o nome esatto del caso sorgente"),
        ("Behavior", "0 = Steel; 1 = Long term con n ∞; 2 = Short term con n0"),
        ("Beam IDs", "Facoltativo; vuoto significa tutti gli elementi applicabili"),
    ], [1.4, 5.54])
    p("Creare G1 con Case = G1_GETTO e Behavior = 0; G2 con G2_PERMANENTI e Behavior = 1; Q con Q_TRAFFICO e Behavior = 2. La lista Beam IDs limita l’applicazione: usare gli ID effettivi, non l’indice della posizione nella lista Grasshopper.")
    h("Bridge Shrinkage Action per deformazioni imposte")
    table(["Ingresso", "Impostazione"], [
        ("Action ID", "RITIRO, o una chiave distinta per ogni incremento"),
        ("Model e Case", "Modello n ritiro analizzato e RITIRO_FEM per gli effetti secondari"),
        ("Strain increment", "Strain del calcestruzzo adimensionale, finito e non positivo"),
        ("Beam IDs", "Facoltativo; seleziona gli elementi interessati"),
        ("Primary Only", "false normalmente; true esclude il secondario FEM"),
        ("Response", "0 = n ritiro; 1 = n0 per deformazione istantanea imposta"),
    ], [1.4, 5.54])
    p("Model e Case sono opzionali nell’interfaccia solo per consentire Primary Only. Con Primary Only = false devono essere presenti. Response = 1 richiede proprietà e azioni equivalenti coerenti con n0: cambiare il selettore non ricalcola automaticamente il modello Straus.")
    p("Le uscite Action entrano nelle rispettive Bridge Phase. Non riutilizzare lo stesso Action ID per contributi diversi. Se il caso non è ancora stato analizzato, completare prima l’analisi: rinominare il caso nel pannello non crea risultati mancanti.")

    # Page 9
    page(8, "Numerare le fasi e assemblare il workflow")
    h("Bridge Phase")
    table(["Ingresso", "Significato e vincoli"], [
        ("Number", "Intero non negativo e univoco; determina l’ordine cronologico"),
        ("Name", "Descrizione leggibile della fase"),
        ("Section State", "0 = Steel; 1 = Composite; 2 = ConcreteExcluded"),
        ("Actions", "Lista delle azioni della fase; può essere vuota per una sola attivazione"),
        ("Day", "Facoltativo, non negativo; metadato senza evoluzione automatica nel tempo"),
    ], [1.4, 5.54])
    steps([
        "Creare la fase 10, ad esempio Acciaio e getto, con Section State = 0 e l’azione G1.",
        "Creare le fasi 20, 30 e 40 con Section State = 1, collegando rispettivamente G2, Q e RITIRO. Usare numeri distanziati permette di inserire successivamente fasi intermedie.",
        "Raccogliere le uscite Phase in una lista. L’ordinamento è dato da Number, non dalla posizione dei fili. Rinumerare una fase non cambia il nome del caso Straus né l’Action ID.",
    ])
    p("Section State descrive lo stato della sezione; Behavior descrive la risposta dell’azione. Dopo l’attivazione della soletta non si può tornare allo stato Steel. ConcreteExcluded richiede un’attivazione precedente e il metodo LinearHistory.")
    h("Build BridGHe Model")
    table(["Ingresso corrente", "Collegamento"], [
        ("Models", "Lista dei modelli di riferimento, con geometria e identificativi coerenti"),
        ("Phases", "Lista delle Bridge Phase appena create"),
        ("Combinations", "Albero della matrice descritto nel capitolo 9"),
        ("Calculation", "0 = Cumulative; 1 = LinearHistory"),
    ], [1.4, 5.54])
    p("Ogni azione conserva già il proprio modello analizzato e il caso sorgente. Models fornisce i riferimenti geometrici e non assegna il significato dei carichi attraverso l’ordine della lista. Avviare Build con Run e leggere le uscite Phases e Diagnostics prima di collegare il verificatore.")
    p("**L’ultimo ingresso è Calculation:** collegarvi un caso di carico produce un errore di tipo. I vecchi collegamenti del merge devono essere ricostruiti con Action e Phase; non è prevista compatibilità con la precedente interfaccia.")

    # Page 10
    page(9, "Preparare combinazioni e metodo di calcolo")
    p("Combinations è un albero Grasshopper con **una riga per ramo**. Il primo ramo contiene l’intestazione; le colonne successive alla prima sono gli Action ID esatti. Ogni riga di dati contiene il nome della combinazione e i coefficienti nello stesso ordine.")
    table(["Ramo", "Nome", "G1", "G2", "Q", "RITIRO"], [
        ("{0}", "vuoto", "G1", "G2", "Q", "RITIRO"),
        ("{1}", "SLS_DEMO", "1", "1", "1", "1"),
        ("{2}", "TEST_Q20", "1.3", "1.5", "20", "1"),
        ("{3}", "TEST_G2_10", "1", "10", "1.3", "1"),
    ], [.58, 2.32, .70, .70, .70, 1.94], centers=(0, 2, 3, 4, 5))
    p("La tabella riproduce REV003. SLS_DEMO è una combinazione dimostrativa; le altre due sono prove con fattori maggiorati. **Non sono combinazioni normative.** Non appiattire tutta la matrice in una sola lista e non eliminare la prima cella del ramo di intestazione.")
    p("Per costruirla manualmente, raccogliere i valori di ciascuna riga e assegnare un percorso distinto con i componenti di gestione degli alberi. Un foglio Excel può fornire gli stessi dati, purché conservi questa struttura. I coefficienti devono essere numerici e finiti.")
    p("La colonna RITIRO scala insieme gli effetti secondari del caso FEM e lo strain primario. Un coefficiente 0.7 produce 0.7 × entrambi i contributi. Non creare colonne manuali chiamate RITIRO [FEM] o RITIRO [strain]: sono nomi interni generati dal workflow.")
    h("Scegliere Calculation")
    table(["Metodo", "Cosa calcola", "Limite operativo"], [
        ("0\nCumulative", "Metodo di sezione gPC cumulativo; raggruppa gli incrementi per comportamento.", "L’ordine è tracciato, ma non introduce un algoritmo generale di attivazione."),
        ("1\nLinearHistory", "Storia lineare N–Mx e attivazione della soletta con HBridgeHistoryResults.Calculate.", "Non esegue il FEM globale per fasi; taglio, pioli e irrigiditori non sono disponibili."),
    ], [1.18, 2.91, 2.85])
    p("Usare Cumulative per riprodurre il percorso completo di REV003 e i controlli di dettaglio compatibili. Usare LinearHistory per studiare esplicitamente la storia lineare di sezione e l’eventuale ConcreteExcluded, mantenendo analisi Straus coerenti per ogni incremento.")
    p("In LinearHistory la deformazione è riferita geometricamente all’inizio della storia; le tensioni del calcestruzzo seguono la sua attivazione. Day non integra viscosità, maturazione o ritiro nel tempo. L’eventuale tau = NaN indica un risultato non disponibile, non nullo.")

    # Page 11
    page(10, "Assemblare ed eseguire il verificatore")
    steps([
        "Selezionare i beam della trave da verificare, con **Super Element Filter** o i relativi attributi. In REV003 la selezione è la spina centrale G2.",
        "Passare la selezione a **Composite beam**. Collegare Composite Section a Section Start; Section End è facoltativo e, se omesso, usa la sezione iniziale. Per una sezione variabile assegnare anche quella finale.",
        "Collegare il Model di Build BridGHe Model a **Super Element Check → BridGHe Model** e la lista delle Composite beam a Super-Element Beam.",
        "Impostare Web Local Instability, Top Flange Local Instability e Bottom Flange Local Instability secondo la verifica richiesta. REV003 li disattiva per isolare il comportamento dimostrativo delle fasi: non assumere che l’instabilità sia già verificata.",
        "Premere Run sul verificatore. Leggere messaggi e risultati prima di attivare i grafici o i dettagli.",
    ])
    table(["Uscita", "Contenuto e uso"], [
        ("Tension Check", "Albero dei risultati di sezione; ingresso Results di Bridge Stress Plot e risultati puntuali per i dettagli compatibili"),
        ("Tau b", "Uscita di taglio non disponibile nel percorso attuale; NaN non equivale a verifica soddisfatta"),
        ("Tension Forces", "Albero delle azioni; ingresso Results di Bridge Force Plot"),
    ], [1.35, 5.59])
    p("La struttura dei risultati segue beam e stazioni; le combinazioni sono nell’ordine del modello. Conservare l’albero **{beam;station}** e la stessa lista ordinata di Composite beam fino ai grafici. Un Flatten indiscriminato può perdere l’associazione fisica dei risultati.")
    h("Campo della verifica")
    p("Il checker H gestisce il campo di azioni implementato e rifiuta componenti fuori campo, fra cui torsione e flessione sull’asse minore. Le azioni FEM non supportate non vengono azzerate per ottenere un esito. Per travi laterali, ponti curvi o azioni eccentriche occorrono verifiche dedicate.")
    p("Le azioni di sezione devono essere riferite coerentemente al baricentro lordo previsto dalla fase; il trasferimento interno fra sezione lorda ed efficace non corregge automaticamente un riferimento FEM eccentrico arbitrario.")
    p("Per l’esempio centrale simmetrico, confrontare prima segni e valori dei singoli contributi G1, G2, Q e RITIRO, poi la combinazione. Una verifica delle tensioni non copre automaticamente stabilità globale, fatica, connessioni, soletta locale o tutte le combinazioni di progetto.")

    # Page 12
    page(11, "Aggiungere pioli e irrigiditori")
    p("I componenti di dettaglio completano controlli locali sui risultati compatibili. Non aggiungono automaticamente massa, rigidezza o collegamenti al modello Straus. Se un dettaglio influenza il FEM, occorre rappresentarne separatamente l’effetto nel modello.")
    steps([
        "Selezionare un risultato puntuale di Tension Check, mantenendo identificabili beam, stazione e combinazione. Collegarlo a **Bridge Detail Check**.",
        "Creare **Bridge Studs** e collegarlo all’ingresso Studs. Specificare il numero di pioli per fila, geometria, passi e materiale; non accettare i valori predefiniti come dati di progetto.",
        "Se necessari, creare **Bridge Intermediate Stiffener** e **Bridge Support Stiffener** e collegarli ai rispettivi ingressi. Inserire le azioni di compressione e le reazioni pertinenti alla sezione e alla combinazione considerate.",
        "Collegare le opzioni di verifica e gli eventuali dettagli aggiuntivi, come armature trasversali o dettagli di fatica, solo quando coerenti con il caso di progetto.",
        "Leggere insieme Check, Demand, Resistance, Units, Ratio, Status e Warnings. Un valore NaN o un controllo non disponibile non costituisce un esito positivo.",
    ])
    table(["Dato", "Unità e controllo"], [
        ("CountPerRow", "Numero intero di pioli per fila, da 1 a 20"),
        ("Diameter, Height e passi", "mm; LongitudinalPitch lungo la trave, TransversePitch fra pioli"),
        ("Fu e GammaV", "Fu in MPa; coefficiente GammaV esplicito e coerente con il criterio adottato"),
        ("Testa e copriferro", "HeadDiameter, HeadThickness e RequiredCover in mm"),
        ("ReactionKN e ExternalCompressionKN", "kN, anche se il modello generale usa N"),
        ("RepeatedActionDetail", "Opzione di dettaglio; non attiva da sola una verifica completa a fatica"),
    ], [2.55, 4.39])
    h("Scelta dei risultati e stato limite")
    p("Usare la combinazione che governa il controllo specifico: il massimo momento non coincide necessariamente con massimo taglio, massima reazione o massima domanda sui pioli. Conservare la provenienza del risultato scelto ed esaminare l’inviluppo delle combinazioni pertinenti.")
    p("Cambiare lo stato limite nelle opzioni non maggiora le azioni. Se il risultato è di esercizio, rimane tale finché non viene calcolata una combinazione coerente con il controllo richiesto. Coefficienti e resistenze devono provenire dal criterio di progetto adottato.")
    p("**LinearHistory non è accettato da Bridge Detail Check** per questi controlli. Usare un percorso numerico compatibile e verificato; non riclassificare un risultato storico come cumulativo per aggirare il messaggio.")
    p("Per i diagrammi collegare i Section result aggiornati a **Bridge Shear Plot** o **Bridge Stud Plot**, mantenendo l’albero beam–stazione e l’ordine delle combinazioni. Le uscite SVG si visualizzano come descritto nel capitolo 12; per i pioli sono disponibili flusso, forza, resistenza e rapporto domanda/resistenza.")

    # Page 13
    page(12, "Visualizzare modello e grafici in Grasshopper")
    h("Preview del modello")
    p("Collegare **Display Settings → Preview BridGHe Model**. Impostare geometria, sezioni, armature, carichi, appoggi, svincoli, assi locali ed etichette. Regolare Symbol scale, Load scale e Transparency; Load case filtra per nome esatto. Il menu BridGHe → Visualizzazione e BridGHe Panel gestiscono le preferenze globali.")
    p("Se si usa una preview dedicata, disattivare le preview automatiche sovrapposte a monte. L’interruttore generale continua a controllare la visualizzazione. Controllare le sezioni composte anche dopo gli override: l’omogeneizzazione numerica non deve far perdere la descrizione fisica.")
    h("Collegamenti dei plot")
    table(["Ingresso", "Bridge Force Plot", "Bridge Stress Plot"], [
        ("BridGHe Model", "Model del Build", "Lo stesso Model"),
        ("Super-Element Beam", "Lista delle Composite beam", "La stessa lista ordinata"),
        ("Results", "Tension Forces", "Tension Check"),
        ("Quantity", "Azione da rappresentare", "Tensione o deformazione"),
        ("Combinations", "Selezione facoltativa delle curve", "Selezione facoltativa delle curve"),
    ], [1.38, 2.78, 2.78])
    steps([
        "Conservare gli alberi dei risultati. Selezionare Quantity: per le azioni 0 = N, 1 = Vx, 2 = Vy, 3 = T, 4 = Mx, 5 = My. I plot convertono forze in kN e momenti in kNm; le tensioni restano in MPa.",
        "Impostare Graphic Title e premere Draw, oppure usare Keep Running = true. File Path è facoltativo; un percorso .svg o .png abilita il salvataggio nel formato corrispondente.",
        "Collegare **SVG → Bridge Plot Preview → SVG**. Index seleziona il grafico da una lista, partendo da 0. Width regola la larghezza da 240 a 1600 pixel; il valore iniziale è 600.",
        "Fare doppio clic sul grafico o scegliere Apri grafico per la finestra dedicata. La preview segue gli aggiornamenti e consente anche la consultazione al 100% o 200%.",
    ])
    h("Come leggere la scala e gli estremi")
    p("Le tacche dell’asse verticale sono ancorate allo zero, quando visibile, con passi della serie 1–2–5 per potenze di dieci: per esempio 5, 10, 20 o 50. Un massimo di 14.754 può quindi avere tacche 0, 5, 10 e 15. I dati numerici e gli estremi non vengono arrotondati dalla scelta delle tacche.")
    p("La fascia grigia è l’inviluppo di tutte le combinazioni, anche se si mostrano poche curve. Le uscite Minimum e Maximum riguardano le curve selezionate e possono differire dagli estremi annotati sull’inviluppo. I tratti collegano i campioni disponibili senza smoothing e senza colmare lacune o salti fra elementi.")

    # Page 14
    page(13, "Preparare e provare i carichi mobili")
    steps([
        "Definire i carichi elementari con **Moving Point Force** o **Moving Distributed Force**, controllando coordinate, intensità e riferimento degli assi.",
        "Raccoglierli in **Moving Vehicle**. Impostare direzione e parametri del veicolo; Lanes usa indici da 1 e, se vuoto, considera tutte le corsie applicabili.",
        "Creare **Moving Load Template** con veicoli e fattori di corsia. Definire tolleranze e interazione con il modello secondo la geometria effettiva.",
        "Creare **Moving Load Path**: Start, End e Lateral devono definire il riferimento del percorso con tre punti non allineati. Impostare Target, Shape, Surface e Divisions; attivare Generate static cases.",
        "Collegare i percorsi a **Attach Moving Paths** insieme al modello. Passare il risultato a **Moving Static Analysis**, assegnare Output path e attivare Run quando la licenza è disponibile.",
        "Leggere Model e Cases. Collegare il modello analizzato e il caso della posizione scelta a Bridge Action. Le posizioni alternative dello stesso veicolo non devono essere sommate indiscriminatamente nella stessa combinazione.",
    ])
    p("Questo percorso genera e analizza posizioni statiche. Non è un’analisi dinamica né un inviluppo continuo esatto. Aumentare Divisions e verificare la sensibilità alla discretizzazione del percorso e della mesh. Per percorsi curvi, la preview dei punti non dimostra da sola la correttezza dell’intera traiettoria nativa.")
    h("Cinque benchmark preparati")
    table(["Prova", "Scopo"], [
        ("01 e 02", "Forza singola con 13 e 25 posizioni; sensibilità al campionamento"),
        ("03", "Percorso invertito; coerenza del verso e delle posizioni"),
        ("04", "Veicolo tandem; contributo di più forze"),
        ("05", "Forza raddoppiata; linearità della risposta"),
    ], [1.1, 5.84])
    p("I benchmark usano una trave semplicemente appoggiata da 6000 mm, E = 210000 MPa e I = 450000000 mm⁴. Confrontano reazioni, momenti alle stazioni disponibili e frecce con soluzioni analitiche. Sono prove del trasferimento e della lettura, non modelli normativi di traffico.")
    code("# Dalla root del repository, preparazione senza Straus\npowershell -ExecutionPolicy Bypass -File tools/test-moving-loads.ps1\n# Esecuzione con licenza disponibile\npowershell -ExecutionPolicy Bypass -File tools/test-moving-loads.ps1 -Native\n# Per una sola prova aggiungere: -Scenario 04")
    p("Stato alla data della guida: cinque scenari preparati e 11 controlli offline superati; la nuova esecuzione nativa resta da effettuare. Sul ponte completo occorrono anche posizionamento trasversale, combinazioni e controlli del campo di validità del verificatore.")

    # Page 15
    page(14, "Eseguire il ponte di esempio REV003")
    p("Aprire **example/19-flusso-acciaio-misto/19-flusso-acciaio-misto_REV003.ghx**. Le revisioni precedenti e i relativi fogli Excel sono in archivio; i nuovi modelli Straus sono in calcoli. La nuova geometria non usa i risultati della precedente trave singola.")
    table(["Elemento", "Dati principali in mm"], [
        ("Spina centrale G2", "H 2000; flangia inferiore 800 × 40; superiore 600 × 30; anima 15"),
        ("Laterali G1 e G3", "H 1600; flangia inferiore 600 × 40; superiore 500 × 30; anima 16"),
        ("Traversi", "H 800; flange 300 × 20; anima 12; passo 5000"),
        ("Soletta per trave", "C32/40; spessore 250; larghezza efficace assegnata 3000"),
        ("Armature B450", "Superiore 3000 mm² a 45 mm; inferiore 2000 mm² a 50 mm"),
    ], [1.65, 5.29])
    steps([
        "Verificare 75 nodi e 98 beam, la disposizione degli appoggi del capitolo 3 e le sezioni. La larghezza nominale del ponte è 12 m; la larghezza efficace di 3 m per trave è un’ipotesi assegnata.",
        "Controllare Eₛ = 210000 MPa, n0 = 7, n ∞ = 18, n ritiro = 12 e strain = −0.0002. Le due file di armatura devono essere attive.",
        "Controllare le copie G1, G2, Q e RITIRO e i quattro percorsi .st7. Verificare carichi da 33, 10 e 10 N/mm su ciascuna trave e assenza di un ulteriore peso proprio automatico.",
        "Con licenza disponibile, attivare CALCOLA 4 MODELLI STRAUS, attendere la conclusione e riportarlo a false. Controllare i casi G1_GETTO, G2_PERMANENTI, Q_TRAFFICO e RITIRO_FEM.",
        "Controllare l’equilibrio: ΣRz = 5940 kN per G1 e 1800 kN per G2 e Q, a segni opposti ai carichi. Per il ritiro verificare l’equilibrio globale delle reazioni e la plausibilità della deformata.",
        "REV003 seleziona G2 dai quattro modelli analizzati prima delle azioni e del Build. Premere Run su Build BridGHe Model e verificare fasi 10, 20, 30, 40, quattro Action ID e nessun errore di corrispondenza. Iniziare da Calculation = 0.",
        "Premere Run su Super Element Check per la selezione G2. Aggiornare i due plot e leggere nel visualizzatore la combinazione SLS_DEMO, quindi le due combinazioni di prova.",
        "Salvare la copia GHX insieme ai modelli analizzati, alle combinazioni e ai grafici. Annotare modifiche ai parametri e versione del plugin per rendere il risultato riproducibile.",
    ])
    p("**Limiti dell’esempio:** verifica dimostrativa della spina sotto azioni simmetriche. La soletta trasversale non è modellata e i traversi restano di acciaio. Restano da completare combinazioni di progetto, traffico mobile sul graticcio, azioni orizzontali, fatica, connessione, instabilità, soletta locale e sensibilità alla mesh. Le componenti fuori campo delle travi laterali richiedono un percorso dedicato.")

    # Page 16
    page(15, "Diagnosticare gli errori e aggiornare i risultati")
    table(["Sintomo", "Causa da controllare", "Intervento"], [
        ("Build rosso collegando un caso in coda", "L’ultimo ingresso è Calculation", "Collegare 0 o 1; inserire il caso in Bridge Action o Bridge Shrinkage Action"),
        ("Caso non trovato", "Nome diverso o risultati non letti", "Usare il nome esatto e rieseguire l’analisi includendo quel caso"),
        ("Risultati mancanti", "Modello non analizzato o campioni incompleti", "Controllare Model in uscita da Straus e le stazioni di ogni beam"),
        ("ID duplicati o geometrie incoerenti", "Copie ricostruite con numerazione o assi diversi", "Ripartire dal modello comune e conservare ID e orientamenti"),
        ("Valori duplicati nelle combinazioni", "Azioni cumulative passate come incrementi", "Separare i contributi e controllare ogni caso prima della somma"),
        ("Ritiro troppo grande", "Strain in µε anziché adimensionale o doppio primario", "Usare −0.0002 per −200 µε; rimuovere la duplicazione tramite Composite Beam Strain"),
        ("Plot vuoto o incoerente", "Albero appiattito, uscita errata o dati non aggiornati", "Forces da Tension Forces; stresses da Tension Check; mantenere gli alberi e premere Draw"),
        ("Tau NaN o dettagli rifiutati", "Metodo o grandezza non supportati", "Non interpretare NaN come zero; usare un percorso di verifica compatibile"),
        ("Errore per torsione o asse minore", "Azioni fuori dal campo del checker H", "Mantenere le azioni FEM e verificarle con un metodo dedicato"),
        ("DLL non copiabile dopo build", "Rhino mantiene il file caricato", "Salvare, chiudere Rhino, ricompilare e riavviare"),
    ], [1.73, 2.14, 3.07])
    h("Ordine di aggiornamento")
    p("Dopo una modifica seguire sempre la catena: **modello e sezioni → analisi delle fasi interessate → Action → Phase → Run del Build → Run del Check → Draw dei plot**. Non confrontare grafici derivati da analisi di versioni diverse dello stesso modello.")
    h("Dati da conservare con ogni esecuzione")
    p("Registrare versione del plugin e delle librerie, geometria, ID, unità, sezioni, rapporti n, casi sorgente, numeri di fase, strain incrementali, matrice delle combinazioni, metodo Calculation e messaggi diagnostici. Archiviare almeno un controllo di equilibrio e un confronto numerico indipendente significativo.")

    # Page 17
    page(16, "Prove disponibili e riferimenti")
    p("La Release con gli esempi 33–45 supera i controlli offline. Per i ponti 11–18 si riportano replay già registrati. Questo aggiornamento non include nuove analisi Straus.")
    table(["Ambito", "Verifica eseguita", "Stato"], [
        ("Release", "Compilazione con 0 warning e 0 errori", "Eseguita"),
        ("Smoke e regressioni", "218 controlli; 54 gruppi principali e 2 grafici", "Superati"),
        ("Conservazione numerica", "238 valori di baseline confrontati senza differenze", "Verificata"),
        ("Esempi 11–18", "344 confronti usando risultati nativi archiviati", "Replay superato; nessuna nuova analisi"),
        ("Esempio 19 REV003", "15 controlli offline; generazione e riapertura del grafo", "Nuovo calcolo Straus da eseguire"),
        ("Carichi mobili", "5 scenari e 11 controlli offline", "Benchmark nativi da eseguire"),
    ], [1.55, 3.43, 1.96])
    h("Comandi dalla root del repository")
    code("# Release e suite senza nuove analisi Straus\npowershell -ExecutionPolicy Bypass -File tools/build.ps1 -Smoke\n\n# Preparare il progetto di test\ndotnet build tests/BridGHe.Smoke -c Release\n\n# Esempio 19 offline\ndotnet run --project tests/BridGHe.Smoke -c Release --no-build -- `\n  --bridge-workflow\n\n# Esempio 19 con licenza e controlli nativi\ndotnet run --project tests/BridGHe.Smoke -c Release --no-build -- `\n  --bridge-workflow --native")
    p("I comandi sono in sintassi PowerShell. Il carattere ` a fine riga continua il comando. Gli script rigenerano i file degli esempi: conservare prima una copia delle modifiche manuali. Per i carichi mobili usare i comandi del capitolo 13.")
    h("Riferimenti locali verificabili")
    p("**Workflow e interfacce:** docs/FASI-E-AZIONI.md; docs/COMPONENTI.md; migration/current-components.json.\n**Grafici:** docs/GRAFICI.md.\n**Esempio corrente:** example/19-flusso-acciaio-misto/README.md e validation-REV003.json nella stessa cartella.\n**Stato complessivo:** example/00-documentazione/validazioni/numbered-phases.json.\n**Riproduzione carichi mobili:** tools/test-moving-loads.ps1.")
    p("Per ricostruire i dettagli implementativi: BridgeWorkflowComponents.cs e BuildBridgeModel.cs in src/BridGHe.Grasshopper/Components/BridgeChecks. Questa guida non seleziona prescrizioni normative o coefficienti di progetto: documenta il trasferimento dei dati e i limiti del calcolo disponibile.")
    page(17, "Controllare il modello e preparare una fase")
    p("L’esempio **23 audit modello fase** introduce Bridge Model Audit e Bridge Phase Model. I due componenti si collocano tra Build Model e l’analisi Straus. Il primo rende leggibili alcuni errori di input; il secondo crea una copia indipendente per il nuovo incremento.")
    h("Bridge Model Audit")
    steps([
        "Collegare il modello FEM a Model. Impostare Tolerance nelle unità di lunghezza del modello: in N–mm il valore è in mm. La tolleranza serve a riconoscere vertici coincidenti, non a correggere automaticamente la geometria.",
        "Leggere Severity, Code, Element e Description come colonne della stessa tabella. Gli errori comprendono identificativi vuoti o duplicati, beam con proprietà/geometria non valide e lunghezze inferiori alla tolleranza. I warning comprendono parti geometricamente disconnesse e nodi isolati.",
        "Usare No input errors come esito dei soli controlli elencati. Un modello con warning può restituire true. La stabilità, i rilasci, il rango della rigidezza, la qualità della mesh e la completezza delle azioni richiedono controlli distinti.",
    ])
    p("La connessione geometrica usa estremi delle beam e vertici delle plate. Due aste che si incrociano internamente senza un nodo comune non risultano connesse. La presenza di un attributo di vincolo non dimostra che tutti i moti rigidi siano impediti.")
    h("Bridge Phase Model")
    table(["Ingresso", "Uso operativo"], [
        ("Model e Name", "Modello comune in acciaio e nome della nuova analisi"),
        ("Beam IDs", "Selezione esplicita delle travi da rendere miste; non è ammessa una lista vuota"),
        ("Property", "Beam Property della risposta n0, n infinito o n ritiro; unità coerenti con Model"),
        ("Clear loads", "True: azzera carichi, casi e percorsi mobili. False: conserva i carichi del modello sorgente"),
    ], [1.4, 5.54])
    p("La copia mantiene geometria, ID, vincoli e altri attributi. I risultati precedenti, le combinazioni e le ricette di fase vengono sempre cancellati perché riferiti all’analisi originaria. Solo le proprietà delle beam selezionate sono sostituite. La preview usa il nuovo modello e la relativa sezione fisica quando disponibile.")
    p("**Sequenza:** acciaio → Audit → Phase Model → nuovi carichi incrementali → Straus → Bridge Action. Phase Model non trasferisce uno stato di deformazione da un’analisi all’altra e non esegue una costruzione per fasi globale. Nell’esempio 23 le due travi sono volutamente disconnesse: DISCONNECTED è atteso.")

    page(18, "Leggere inviluppi e valori concomitanti")
    p("L’esempio **24 inviluppi concomitanti** mostra Bridge Result Cases e i tre inviluppi: Bridge Beam Envelope, Bridge Bearing Envelope e Bridge Displacement Envelope. Operano sui risultati già memorizzati nel modello; non lanciano Straus e non generano nuove combinazioni.")
    h("Selezionare casi e grandezza")
    steps([
        "Collegare il modello analizzato a Bridge Result Cases. Cases elenca i nomi effettivamente presenti nei risultati; le altre colonne riportano caso, passo, tipo e numero dei record. Un caso definito ma non analizzato non compare.",
        "Collegare Model all’inviluppo richiesto. Element IDs e Cases sono opzionali: vuoti selezionano tutti gli elementi del tipo richiesto e tutti i relativi casi disponibili. Selezionare gli appoggi esplicitamente per le reazioni. Ogni elemento selezionato deve avere risultati.",
        "Impostare Component secondo la tabella e Mode: 0 minimo, 1 massimo, 2 massimo valore assoluto. Mode=2 mantiene il segno del valore governante.",
        "Associare ciascuna riga di Value a Element ID, Station, Case, Step e Step type. Il ramo Concomitant con lo stesso indice contiene le sei componenti del medesimo campione governante.",
    ])
    table(["Componente", "Indici 0 1 2 3 4 5", "Riferimento"], [
        ("Beam Envelope", "N  V2  V3  T  M2  M3", "Assi locali Straus"),
        ("Bearing Envelope", "Fx  Fy  Fz  Mx  My  Mz", "Assi globali"),
        ("Displacement Envelope", "Ux  Uy  Uz  Rx  Ry  Rz", "Assi globali; rotazioni rad"),
    ], [1.85, 2.7, 2.39])
    p("Le unità restano quelle del modello. In N–mm, forze e reazioni sono in N, momenti in Nmm e spostamenti in mm. Station è una distanza dal primo estremo della beam; eventuali stazioni normalizzate sono convertite usando la lunghezza dell’elemento. Per i nodi Station vale zero.")
    h("Concomitanza e copertura")
    p("Se governa M3 del caso Q, tutti gli altri valori di Concomitant provengono da Q allo stesso passo e alla stessa stazione. Non sono assemblati i massimi indipendenti di N, V e M. Step conserva l’identificatore disponibile: non è automaticamente una coordinata del veicolo.")
    p("Coverage segnala i casi mancanti a una stazione. Non sono eseguite interpolazioni tra griglie di stazioni o passi diversi. Campioni duplicati, valori non finiti e selezioni inesistenti causano errore. L’esempio usa dati numerici controllati: Fz governante = −120 N e Uz governante = −10 mm; non rappresenta una nuova analisi strutturale.")

    page(19, "Costruire combinazioni e incrementi di ritiro")
    h("Bridge Phase Table")
    p("Collegare la lista delle Bridge Phase a Phases. L’uscita Phases restituisce gli oggetti ordinati per Number. Le colonne mostrano numero, nome, giorno, stato della sezione, Action ID, comportamento, modello sorgente, caso e strain. Una fase di sola attivazione rimane visibile con Action ID vuoto.")
    p("Il componente controlla unicità di numeri e identità, Action ID non duplicati e giorni non decrescenti quando assegnati. La corrispondenza effettiva di casi, geometrie e risultati resta verificata da Build BridGHe Model. Il giorno è un dato di sequenza, non un’integrazione automatica della viscosità.")
    h("Bridge Action Combinations")
    steps([
        "Collegare le fasi ordinate a Phases. Scrivere in Action IDs gli identificativi delle colonne che si vogliono assegnare, nell’ordine scelto.",
        "Fornire una lista Names e un albero Coefficients: un ramo per combinazione, un numero per Action ID. L’ordine dei nomi corrisponde all’ordine dei percorsi dei rami in Grasshopper.",
        "Collegare Table direttamente all’ingresso Combinations del Build. La prima riga contiene gli Action ID; le successive nomi e fattori. Non appiattire questo albero.",
    ])
    p("Le azioni omesse ricevono zero e sono elencate in Omitted actions. I fattori devono essere finiti; nomi duplicati, ID sconosciuti e rami di lunghezza errata vengono rifiutati. La scelta normativa dei coefficienti e delle simultaneità rimane esplicita a carico dell’utente.")
    p("Nell’**esempio 25**, le colonne CS,G1 e i valori 0.7,1 diventano una tabella ordinata G1,CS con 1,0.7. Il fattore CS scala l’intera azione di ritiro. L’esempio attiva Primary Only e quindi non comprende le forze secondarie da vincolo.")
    h("Bridge Shrinkage Increments")
    p("Nell’**esempio 27**, Days = 7,28,100 e Cumulative strain = −0.0001,−0.0002,−0.00025 producono Increment = −0.0001,−0.0001,−0.00005, con Initial strain = 0. Ogni valore va alla Bridge Shrinkage Action della stessa posizione temporale.")
    p("La regola è Δεᵢ = εᵢ − εᵢ₋₁; per il primo valore si usa Initial strain. La somma deve coincidere con l’ultimo valore cumulativo meno quello iniziale. I giorni devono crescere strettamente; il ritiro cumulativo deve essere finito, negativo o nullo e non crescente. Un recupero di deformazione richiede un diverso modello di azione.")
    p("Gli strain sono adimensionali: −200 µε corrispondono a −0.0002. Il componente converte una storia fornita dall’utente; non calcola una legge normativa di ritiro o viscosità e non modifica i risultati FEM.")

    page(20, "Riepilogare le verifiche e confrontare analisi")
    h("Bridge Check Summary")
    p("L’**esempio 26** collega Bridge Section Check al nuovo riepilogo. Results accetta i BridgeStressResult prodotti da Section Check, Super Element Check e Detail Check. Labels richiede una voce per risultato, per esempio G2/x=15000/SLS, in modo da mantenere tracciabili sezione e combinazione.")
    table(["Uscita", "Interpretazione"], [
        ("Label Phase Check", "Identificano ogni controllo locale estratto dai risultati GPC"),
        ("Ratio Status Units Notes", "Rapporto, stato, unità e note originali; le grandezze non vengono ricalcolate"),
        ("Maximum finite ratio", "Massimo tra i rapporti numerici presenti; NaN se non ne esistono"),
        ("Unavailable ratios", "Numero dei rapporti mancanti o non finiti da interpretare tramite stato e note"),
        ("Governing labels", "Tutte le righe che raggiungono il massimo rapporto finito"),
    ], [2.1, 4.84])
    p("Una verifica non eseguita non vale zero. Anche con Maximum finite ratio inferiore a uno possono esistere controlli mancanti, fuori campo o da svolgere separatamente. Il riepilogo espone i controlli locali disponibili di taglio, connessione e torsione; non produce da solo un giudizio complessivo sul ponte.")
    h("Bridge Result Comparison")
    p("Per confrontare revisioni o discretizzazioni, campionare prima le stesse grandezze nelle stesse posizioni fisiche. Reference keys e Candidate keys devono contenere chiavi univoche uguali, anche in ordine diverso; i due elenchi Values devono avere le stesse unità.")
    p("Il criterio applicato a ogni chiave è **|c − r| ≤ Abs + Rel · |r|**, con r riferimento e c candidato. Abs ha le unità della grandezza; Rel è adimensionale. Per grandezze con unità diverse usare componenti separati o tolleranze opportunamente normalizzate.")
    p("Difference = c − r. Relative difference = (c − r)/|r|; se r = 0 e la differenza non è nulla, l’uscita relativa è NaN e il confronto usa la tolleranza assoluta. Chiavi mancanti, duplicati e numeri non finiti causano errore.")
    p("Nell’**esempio 27**, gli spostamenti −10 e −20 mm vengono confrontati con −10.02 e −20.1 mm, presentati in ordine invertito. Entrambi passano con Rel = 0.01 e Abs = 0.000001 mm. Questo controllo non genera mesh né dimostra da solo la convergenza: per uno studio occorrono più raffinamenti e grandezze fisiche confrontabili.")

    page(21, "Preparare le corsie e riprodurre gli esempi")
    h("Bridge Moving Lanes")
    p("L’**esempio 28** prepara tre percorsi rettilinei a partire da un unico Moving Load Path. Collegare il prototipo a Prototype, gli scostamenti trasversali a Offsets e un nome di caso univoco per ogni percorso a Cases.")
    steps([
        "Impostare nel prototipo veicolo, template, target, numero di divisioni e conversione in casi statici. Questi dati vengono copiati indipendentemente in ciascun percorso.",
        "Fornire offsets distinti e finiti nelle unità del modello. Il verso positivo è quello del terzo punto del prototipo, proiettato perpendicolarmente all’asse rettilineo.",
        "Collegare Paths ad Attach Moving Paths e controllare la preview. Procedere con Moving Static Analysis sul modello desiderato; la creazione dei percorsi non avvia il solutore.",
    ])
    p("Il componente genera percorsi indipendenti, non coefficienti di corsia o simultaneità normative. Per percorsi curvi usare le definizioni individuali di Moving Load Path: una semplice traslazione non crea automaticamente corsie concentriche. L’esempio assegna un veicolo dimostrativo da 100 kN su 20 m, con offsets −3, 0 e +3 m.")
    h("Cartelle e verifiche riproducibili")
    table(["Esempio", "Componenti e controlli"], [
        ("23 audit modello fase", "Diagnostica e copia con invalidazione dei risultati"),
        ("24 inviluppi concomitanti", "Catalogo casi, forze, reazioni e spostamenti"),
        ("25 fasi combinazioni", "Ordine esplicito e matrice per Action ID"),
        ("26 riepilogo verifiche gpc", "Rapporti e stati ricavati dal calcolo GPC"),
        ("27 ritiro confronti", "Bilancio degli incrementi e confronto per chiavi fisiche"),
        ("28 corsie mobili", "Copie indipendenti dei percorsi e preview"),
    ], [2.08, 4.86])
    code("# Controlli dei servizi\ndotnet run --project tests/BridGHe.Smoke -c Release --no-build -- `\n  --regression bridge-utilities\n# Rigenera gli esempi 23–28 e ne verifica salvataggio/riapertura\ndotnet run --project tests/BridGHe.Smoke -c Release --no-build -- `\n  --utility-examples")
    p("Le prove dei nuovi strumenti comprendono 44 controlli di servizio e sei grafi calcolati e riaperti con i 12 nuovi componenti. Non richiedono una nuova analisi Straus. Gli esempi di inviluppo usano campioni controllati, mentre il riepilogo ricalcola la sezione con GPC. Salvare una copia prima di rigenerare esempi modificati manualmente.")
    from guide_bridge_design import append_design_chapters
    append_design_chapters(page, h, p, table, steps, code)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make_document()
