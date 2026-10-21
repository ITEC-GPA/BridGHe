"""Create the editable Word template for example 48. Uses the bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

root = Path(__file__).resolve().parents[1]
folder = root / "example/48-report-word"
folder.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.75)
section.header_distance = section.footer_distance = Inches(.3)
for name, size in [("Normal",11),("Title",24),("Subtitle",12),("Heading 1",16),("Heading 2",12),("Caption",10)]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.underline = False
    for border in style._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for fonts in style._element.xpath('.//w:rFonts'):
        for attr in list(fonts.attrib):
            if attr.lower().endswith('theme'):
                del fonts.attrib[attr]
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.1
    if name.startswith("Heading") or name == "Caption":
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
table_style = doc.styles.add_style("BGTableText", WD_STYLE_TYPE.PARAGRAPH)
table_style.base_style = doc.styles["Normal"]
table_style.font.size = Pt(10)
table_style.paragraph_format.line_spacing = 1.0
table_style.paragraph_format.space_after = Pt(3)
header = section.header.paragraphs[0]
header.add_run("BridGHe    ")
header.add_run("{{BG.Pro")  # Word routinely splits marker text across runs.
header.add_run("ject}}")
header.style = doc.styles["Caption"]
footer = section.footer.paragraphs[0]
footer.add_run("Revisione {{BG.Revision}}    |    ")
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.add_paragraph("{{BG.Title}}", "Title")
doc.add_paragraph("{{BG.Project}}", "Subtitle")
doc.add_paragraph("Autore {{BG.Author}}\nRevisione {{BG.Revision}}\nEsportazione {{BG.Date}}")
doc.add_paragraph("Risultati delle verifiche delle sezioni miste e dei dettagli del ponte. Il riepilogo distingue esiti disponibili, superamenti dei limiti e informazioni mancanti.")
for chapter in ["Summary","Model","Sections","Phases","Checks","Critical","Plots","Appendix"]:
    doc.add_paragraph("{{BG." + chapter + "}}")
doc.core_properties.title = "Template del report di verifica BridGHe"
doc.core_properties.author = "BridGHe"
doc.save(folder / "Template_BridGHe.docx")
print(folder / "Template_BridGHe.docx")
